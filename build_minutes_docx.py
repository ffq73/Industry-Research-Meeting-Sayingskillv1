#!/usr/bin/env python3
"""Build a formatted meeting-minutes DOCX from plain text."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


DEFAULT_WATERMARK = "内部资料 严禁外传 追究责任"


def clear_body(document: Document) -> None:
    body = document._body._element
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_run_font(run, size: float = 10.5, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "楷体_GB2312")
    run.font.size = Pt(size)
    run.bold = bold


def add_paragraph(document: Document, text: str, *, bold: bool = False) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.2
    run = paragraph.add_run(text)
    set_run_font(run, bold=bold)


def add_header_watermark(document: Document) -> None:
    for section in document.sections:
        paragraph = section.header.paragraphs[0] if section.header.paragraphs else section.header.add_paragraph()
        paragraph.text = DEFAULT_WATERMARK
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            set_run_font(run, size=18, bold=True)
            run.font.color.rgb = RGBColor(210, 210, 210)


def set_page_layout(document: Document) -> None:
    for section in document.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.17)
        section.right_margin = Cm(3.17)


def build_docx(title: str, content: str, output: Path, template: Path | None) -> None:
    if template and template.exists():
        document = Document(str(template))
        clear_body(document)
    else:
        document = Document()
        add_header_watermark(document)
    set_page_layout(document)

    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_paragraph.add_run(title)
    set_run_font(title_run, size=10.5, bold=True)

    document.add_paragraph()
    body = content.strip()
    if body.startswith(title):
        body = body[len(title) :].strip()

    for raw_line in body.splitlines():
        line = raw_line.strip()
        if line.startswith("#") or line.startswith("```") or (line.startswith("`") and line.endswith("`")):
            continue
        if not line:
            document.add_paragraph()
            continue
        is_question = bool(re.match(r"^Q(?:\d+)?[：:]", line))
        add_paragraph(document, line, bold=is_question)

    document.save(str(output))


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--title", required=True, help="Document title, e.g. 20251107 恒逸石化 年度策略会线上交流")
    parser.add_argument("--content", required=True, type=Path, help="Plain-text minutes body")
    parser.add_argument("--output", required=True, type=Path, help="Output .docx path")
    parser.add_argument(
        "--template",
        type=Path,
        default=Path(__file__).resolve().parents[1] / "assets" / "meeting-minutes-template.docx",
        help="Optional DOCX template with watermark and styles",
    )
    args = parser.parse_args()

    content = args.content.read_text(encoding="utf-8")
    build_docx(args.title, content, args.output, args.template)
    print(f"Wrote {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
