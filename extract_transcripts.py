#!/usr/bin/env python3
"""Extract text from transcript files and combine them into one UTF-8 file."""

from __future__ import annotations

import argparse
from pathlib import Path


TEXT_EXTENSIONS = {".txt", ".md", ".csv"}


def read_text_file(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "cp936"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="replace")


def read_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx is required to extract .docx files") from exc

    doc = Document(str(path))
    chunks: list[str] = []
    chunks.extend(p.text for p in doc.paragraphs if p.text.strip())
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                chunks.append("\t".join(cells))
    return "\n".join(chunks)


def read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("pypdf is required to extract .pdf files") from exc

    reader = PdfReader(str(path))
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        pages.append(f"[Page {index}]\n{text.strip()}")
    return "\n\n".join(pages)


def extract(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in TEXT_EXTENSIONS:
        return read_text_file(path)
    if suffix == ".docx":
        return read_docx(path)
    if suffix == ".pdf":
        return read_pdf(path)
    if suffix in {".m4a", ".mp3", ".wav", ".aac", ".flac"}:
        raise RuntimeError(f"{path.name} is an audio file. Transcribe it before extraction.")
    raise RuntimeError(f"Unsupported file type: {path.suffix}")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("inputs", nargs="+", type=Path, help="Transcript files to combine")
    parser.add_argument("-o", "--output", type=Path, default=Path("combined-transcript.txt"))
    args = parser.parse_args()

    sections = []
    for input_path in args.inputs:
        path = input_path.resolve()
        if not path.exists():
            raise FileNotFoundError(path)
        text = extract(path).strip()
        sections.append(f"===== {path.name} =====\n{text}")

    args.output.write_text("\n\n".join(sections) + "\n", encoding="utf-8")
    print(f"Wrote {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
